import openpyxl

from strings import raw, t


def generate_daily_excel(date_str, stats, top_sales, top_arrivals,
                         price_changes=None, status_changes=None):
    sheets = raw("excel_sheets")  # [Summary, Top arrivals, Top sales, Price changes, Status changes]
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = sheets[0]
    ws.append([t("daily_header", date=date_str).replace("<b>", "").replace("</b>", "")])
    ws.append([])
    ws.append(["Metric", "Count", "Qty"])
    ws.append(["Sales", stats["sales"][0], stats["sales"][1]])
    ws.append(["Arrivals", stats["arrivals"][0], stats["arrivals"][1]])
    ws.append(["Price changes", stats["price_changes"], ""])
    ws.append(["Status changes", stats["status_changes"], ""])

    def top_sheet(name, items):
        s = wb.create_sheet(title=name)
        s.append(["#", "Title", "Publisher", "Qty"])
        for i, it in enumerate(items or [], 1):
            s.append([i, it["title"], it["publisher"], it["qty"]])

    def list_sheet(name, items):
        s = wb.create_sheet(title=name)
        s.append(["Title", "Publisher", "Change"])
        for it in items or []:
            s.append([it["title"], it["publisher"], it.get("detail", "")])

    top_sheet(sheets[1], top_arrivals)
    top_sheet(sheets[2], top_sales)
    list_sheet(sheets[3], price_changes)
    list_sheet(sheets[4], status_changes)

    filename = f"report_{date_str}.xlsx"
    wb.save(filename)
    return filename
